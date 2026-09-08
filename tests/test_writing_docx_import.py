import asyncio
import io
import unittest

import docx
from fastapi import UploadFile

from api.index import (
    DocxImportValidationError,
    docx_quiz_validation_error,
    parse_docx_to_quiz,
    upload_docx,
    upload_docx_batch,
)


def complete_writing_document():
    document = docx.Document()
    document.add_paragraph("[TITLE] Academic Writing Practice")
    document.add_paragraph("[TIME] 60")
    document.add_paragraph("[TYPE] Writing")
    document.add_paragraph("[WRITING_TASK 1]")
    document.add_paragraph("[TASK_TITLE] Academic Writing Task 1")
    document.add_paragraph("[INSTRUCTIONS] You should spend about 20 minutes on this task. Write at least 150 words.")
    document.add_paragraph("[MINUTES] 20")
    document.add_paragraph("[MIN_WORDS] 150")
    document.add_paragraph("[PROMPT]")
    prompt = document.add_paragraph()
    prompt.add_run("The ").bold = True
    prompt.add_run("chart below shows changes in transport use.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Chart source table"
    document.add_paragraph("[MEDIA] https://cdn.example.test/writing-task-1-chart.png")
    document.add_paragraph("[WRITING_TASK 2]")
    document.add_paragraph("[PROMPT]")
    document.add_paragraph("Some people believe governments should invest more in public transport. Discuss both views and give your opinion.")
    return document


def uploaded_document(document, filename="writing.docx"):
    raw = io.BytesIO()
    document.save(raw)
    raw.seek(0)
    return UploadFile(file=raw, filename=filename)


def uploaded_writing_file(filename="writing.docx"):
    return uploaded_document(complete_writing_document(), filename)


class WritingDocxImportTests(unittest.TestCase):
    def test_writing_docx_returns_tasks_without_generic_questions(self):
        quiz = parse_docx_to_quiz(complete_writing_document())

        self.assertEqual(quiz["type"], "Writing")
        self.assertEqual(quiz["title"], "Academic Writing Practice")
        self.assertEqual(quiz["timeLimit"], 60)
        self.assertEqual(quiz["questions"], [])
        self.assertEqual(docx_quiz_validation_error(quiz), None)
        self.assertEqual([task["taskNumber"] for task in quiz["writingTasks"]], [1, 2])
        self.assertEqual(quiz["writingTasks"][0]["mediaUrl"], "https://cdn.example.test/writing-task-1-chart.png")
        self.assertEqual(quiz["writingTasks"][0]["minimumWords"], 150)
        self.assertEqual(quiz["writingTasks"][0]["recommendedMinutes"], 20)
        self.assertIn("<strong>The </strong>", quiz["writingTasks"][0]["prompt"])
        self.assertIn("<table", quiz["writingTasks"][0]["prompt"])

    def test_writing_requires_both_tasks_and_a_prompt_for_each(self):
        document = docx.Document()
        document.add_paragraph("[TYPE] Writing")
        document.add_paragraph("[WRITING_TASK 1]")
        document.add_paragraph("[PROMPT]")
        document.add_paragraph("Describe the chart.")

        error = docx_quiz_validation_error(parse_docx_to_quiz(document))

        self.assertEqual(error, "Writing DOCX phải có cả [WRITING_TASK 1] và [WRITING_TASK 2].")

    def test_writing_media_must_be_a_public_https_url(self):
        document = docx.Document()
        document.add_paragraph("[TYPE] Writing")
        document.add_paragraph("[WRITING_TASK 1]")
        document.add_paragraph("[PROMPT]")
        document.add_paragraph("Describe the chart.")
        document.add_paragraph("[MEDIA] chart.png")

        with self.assertRaisesRegex(DocxImportValidationError, r"public HTTPS URL"):
            parse_docx_to_quiz(document)

    def test_writing_tasks_must_keep_task_one_before_task_two(self):
        document = docx.Document()
        document.add_paragraph("[TYPE] Writing")
        document.add_paragraph("[WRITING_TASK 2]")
        document.add_paragraph("[PROMPT]")
        document.add_paragraph("Discuss the statement.")

        with self.assertRaisesRegex(DocxImportValidationError, r"WRITING_TASK 1.*before.*WRITING_TASK 2"):
            parse_docx_to_quiz(document)

    def test_single_and_batch_routes_accept_questionless_writing(self):
        single = asyncio.run(upload_docx(uploaded_writing_file("single-writing.docx")))
        batch = asyncio.run(upload_docx_batch([
            uploaded_writing_file("first-writing.docx"),
            uploaded_writing_file("second-writing.docx"),
        ]))

        self.assertTrue(single["success"])
        self.assertEqual(single["quiz"]["questions"], [])
        self.assertTrue(batch["success"])
        self.assertEqual([item["filename"] for item in batch["results"]], ["first-writing.docx", "second-writing.docx"])
        self.assertTrue(all(item["success"] for item in batch["results"]))

    def test_batch_keeps_a_valid_writing_file_when_another_is_incomplete(self):
        incomplete = docx.Document()
        incomplete.add_paragraph("[TYPE] Writing")
        incomplete.add_paragraph("[WRITING_TASK 1]")
        incomplete.add_paragraph("[PROMPT]")
        incomplete.add_paragraph("Describe the chart.")

        batch = asyncio.run(upload_docx_batch([
            uploaded_writing_file("valid-writing.docx"),
            uploaded_document(incomplete, "incomplete-writing.docx"),
        ]))

        self.assertFalse(batch["success"])
        self.assertTrue(batch["results"][0]["success"])
        self.assertFalse(batch["results"][1]["success"])
        self.assertIn("[WRITING_TASK 2]", batch["results"][1]["error"])

    def test_reading_parser_still_uses_question_validation(self):
        document = docx.Document()
        document.add_paragraph("[TITLE] Reading Regression")
        document.add_paragraph("[TYPE] Reading")
        document.add_paragraph("[PASSAGE]")
        document.add_paragraph("A short passage.")
        document.add_paragraph("[QUESTIONS]")
        document.add_paragraph("[BLANK]")
        document.add_paragraph("1. The answer is ____.")
        document.add_paragraph("*example")

        quiz = parse_docx_to_quiz(document)

        self.assertEqual(quiz["type"], "Reading")
        self.assertEqual(len(quiz["questions"]), 1)
        self.assertEqual(docx_quiz_validation_error(quiz), None)


if __name__ == "__main__":
    unittest.main()
