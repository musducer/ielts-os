import json
import base64
import tempfile
import threading
import time
import unittest
from pathlib import Path

import docx

from exam_generator.docx_renderer import render_docx
from exam_generator.pipeline import GENERATOR_SYSTEM, ExamGenerationPipeline, PipelineConfig, PipelineResult
from exam_generator.providers import ProviderError, TextProvider
from exam_generator.roundtrip import verify_docx_roundtrip
from exam_generator.schema import CanonicalExam, CanonicalSchemaError, assign_stable_identifiers, validate_canonical_exam
from exam_generator.grounding import ground_exam, source_hash
from exam_generator.media import (InternalMediaStore, MediaAsset, MediaExtraction, MediaOccurrence,
                                  extract_docx_media, resolve_media_bindings, verify_embedded_media_roundtrip)
from exam_generator.validation import critic_review


def reading_payload():
    return {
        "title": "Rich format contract",
        "exam_type": "Reading",
        "time_limit": 60,
        "sections": [{
            "label": "Passage 1",
            "passage": [
                {"runs": [{"text": "Italic source lead", "italic": True}], "alignment": "left"},
                {"runs": [{"text": "Centered bold heading", "bold": True}], "alignment": "center"},
            ],
            "questions": [
                {
                    "question_number": 1,
                    "block_type": "CHOICE",
                    "text": {"runs": [{"text": "Which statement is correct?", "bold": True}]},
                    "instruction": ["Choose the correct answer."],
                    "options": ["first answer", {"runs": [{"text": "second answer", "italic": True}]}],
                    "correct_answers": ["second answer"],
                    "explanation": {
                        "why": "The passage explicitly identifies the second answer, so the other option contradicts the source wording.",
                        "evidence": {"quote": "Centered bold heading", "source": "passage"},
                        "paraphrase_mapping": "centered heading -> second answer",
                        "correct_reason": "The source wording directly identifies the required second answer.",
                        "distractor_reasons": {"first answer": "The first answer is not named by the evidence."},
                        "student_takeaway": "Match the exact source claim before choosing an option.",
                    },
                },
                {
                    "question_number": 2,
                    "block_type": "BLANK",
                    "text": "Complete the phrase.",
                    "instruction": ["Write ONE WORD ONLY."],
                    "correct_answers": ["heading"],
                    "explanation": {
                        "why": "The supplied phrase contains the required word in the unique sentence that defines the completion.",
                        "evidence": {"quote": "Italic source lead", "source": "passage"},
                        "paraphrase_mapping": "source lead -> required completion",
                        "correct_reason": "The unique source phrase gives the completion word directly.",
                        "student_takeaway": "Use the exact wording around a completion gap to check spelling.",
                    },
                },
            ],
        }],
        "writing_tasks": [],
    }


def structure_payload():
    payload = reading_payload()
    for section in payload["sections"]:
        for question in section["questions"]:
            question["correct_answers"] = []
            question.pop("explanation", None)
    return payload


def source_outline(user: str, *, exam_type: str = "Reading", numbers=(1, 2)):
    source = user.split("<SOURCE_DATA>", 1)[1].split("</SOURCE_DATA>", 1)[0]
    blocks = json.loads(source)
    return {
        "exam_type": exam_type,
        "title": "Rich format contract",
        "sections": [{
            "label": "Passage 1",
            "passage_block_ids": [],
            "groups": [{
                "block_ids": [block["block_id"] for block in blocks],
                "question_numbers": list(numbers),
                "task_number": 0,
            }],
        }],
        "metadata_block_ids": [],
    }


def question_from_call(kwargs):
    return json.loads(kwargs["user"].split("<QUESTION>", 1)[1].split("</QUESTION>", 1)[0])


class CompletePipelineProvider(TextProvider):
    def __init__(self):
        self.calls = []
        self.peak = 0
        self.active = 0
        self.lock = threading.Lock()

    def complete(self, **kwargs):
        self.calls.append(kwargs)
        phase = kwargs.get("phase")
        if phase == "structure_outline":
            return json.dumps(source_outline(kwargs["user"]))
        if phase in {"structure_group", "structure_repair"}:
            return json.dumps(structure_payload())
        if phase in {"question", "question_repair", "explanation_repair"}:
            question = question_from_call(kwargs)
            number = question["question_number"]
            with self.lock:
                self.active += 1
                self.peak = max(self.peak, self.active)
            time.sleep(0.01)
            with self.lock:
                self.active -= 1
            expected = reading_payload()["sections"][0]["questions"][number - 1]
            return json.dumps({
                "question_number": number,
                "correct_answers": expected["correct_answers"],
                "explanation": expected["explanation"],
            })
        if phase == "critic":
            return '{"issues": [{"severity":"PASS","code":"CRITIC_PASS","message":"Verified."}]}'
        raise AssertionError(phase)


class ExamGeneratorPipelineTests(unittest.TestCase):
    def test_canonical_schema_and_index_roundtrip(self):
        exam = CanonicalExam.from_dict(reading_payload())
        assign_stable_identifiers(exam, source_hash("Italic source lead Centered bold heading"))
        first_option = exam.sections[0].questions[0].options[0]
        exam.sections[0].questions[0].explanation.distractor_reasons = {
            first_option.id: "The first answer is not named by the evidence."
        }
        self.assertEqual(validate_canonical_exam(exam), [])
        with tempfile.TemporaryDirectory() as directory:
            output = render_docx(exam, Path(directory) / "reading.docx")
            result = verify_docx_roundtrip(exam, output)

            self.assertEqual(result.quiz["questions"][0]["correctAnswer"], 1)
            self.assertEqual(result.quiz["questions"][1]["correctAnswer"], "heading")

    def test_optional_slots_placeholder_does_not_reject_non_map_question(self):
        payload = reading_payload()
        payload["sections"][0]["questions"][0]["slots"] = ["not used outside map drag"]

        exam = CanonicalExam.from_dict(payload)

        self.assertEqual(exam.sections[0].questions[0].slots, {})

    def test_map_drag_slots_remain_fail_closed(self):
        payload = reading_payload()
        question = payload["sections"][0]["questions"][0]
        question["block_type"] = "MAP_DRAG"
        question["slots"] = [{"x": 10, "y": 20}]

        with self.assertRaisesRegex(CanonicalSchemaError, "MAP_DRAG"):
            CanonicalExam.from_dict(payload)

    def test_optional_diagram_placeholder_does_not_reject_non_diagram_question(self):
        payload = reading_payload()
        payload["sections"][0]["questions"][0]["diagram"] = [{"unused": True}]

        exam = CanonicalExam.from_dict(payload)

        self.assertEqual(exam.sections[0].questions[0].diagram, {})

    def test_diagram_geometry_remains_fail_closed(self):
        payload = reading_payload()
        question = payload["sections"][0]["questions"][0]
        question["block_type"] = "DIAGRAM_LABEL"
        question["diagram"] = [{"id": "not-an-object"}]

        with self.assertRaisesRegex(CanonicalSchemaError, "DIAGRAM_LABEL"):
            CanonicalExam.from_dict(payload)

    def test_renderer_preserves_run_styles_and_alignment(self):
        exam = CanonicalExam.from_dict(reading_payload())
        with tempfile.TemporaryDirectory() as directory:
            output = render_docx(exam, Path(directory) / "style.docx")
            document = docx.Document(output)
            italic = next(paragraph for paragraph in document.paragraphs if paragraph.text == "Italic source lead")
            centered = next(paragraph for paragraph in document.paragraphs if paragraph.text == "Centered bold heading")
            choice = next(paragraph for paragraph in document.paragraphs if paragraph.text == "1. Which statement is correct?")

            self.assertTrue(italic.runs[0].italic)
            self.assertEqual(centered.alignment, 1)
            self.assertTrue(centered.runs[0].bold)
            self.assertTrue(choice.runs[-1].bold)

    def test_pipeline_caches_only_parser_verified_result(self):
        with tempfile.TemporaryDirectory() as directory:
            provider = CompletePipelineProvider()
            pipeline = ExamGenerationPipeline(provider, PipelineConfig(state_dir=Path(directory), max_repairs=0, question_workers=1))
            source = "Italic source lead. Centered bold heading."
            first = pipeline.generate(source, {"exam_type": "Reading"})
            second = pipeline.generate(source, {"exam_type": "Reading"})

            self.assertEqual(first.status, "READY_FOR_REVIEW")
            self.assertFalse(first.cached)
            self.assertEqual(second.status, "READY_FOR_REVIEW")
            self.assertTrue(second.cached)
            self.assertEqual(len(provider.calls), 5)  # outline + group + 2 questions + critic; cache adds no calls.

    def test_provider_timeout_is_failed_not_reported_as_parser_contract(self):
        class TimeoutProvider(TextProvider):
            def complete(self, **_kwargs):
                raise ProviderError("Provider request timed out.")

        with tempfile.TemporaryDirectory() as directory:
            pipeline = ExamGenerationPipeline(TimeoutProvider(), PipelineConfig(
                state_dir=Path(directory), max_repairs=0, question_workers=0,
                provider_attempts=1, provider_timeout_seconds=60,
            ))
            result = pipeline.generate("A short reading source.", {"exam_type": "Reading"})

        self.assertEqual(result.status, "FAILED")
        self.assertEqual(result.validation_state, "FAIL")
        self.assertEqual(result.round_trip_state, "NOT_RUN")
        self.assertIn("timed out", result.error.casefold())
        self.assertEqual(result.issues[0]["code"], "PROVIDER_TIMEOUT")

    def test_pipeline_retries_malformed_canonical_json_inside_provider_boundary(self):
        class SequenceProvider(TextProvider):
            def __init__(self):
                self.generate_calls = 0

            def complete(self, **kwargs):
                phase = kwargs.get("phase")
                if phase == "structure_outline":
                    self.generate_calls += 1
                    return '{"title":' if self.generate_calls == 1 else json.dumps(source_outline(kwargs["user"]))
                if phase == "structure_group":
                    return json.dumps(structure_payload())
                if phase == "question":
                    question = question_from_call(kwargs)
                    expected = reading_payload()["sections"][0]["questions"][question["question_number"] - 1]
                    return json.dumps({"question_number": question["question_number"],
                        "correct_answers": expected["correct_answers"], "explanation": expected["explanation"]})
                if phase == "critic":
                    return '{"issues": [{"severity":"PASS","code":"CRITIC_PASS","message":"Verified."}]}'
                raise AssertionError(phase)

        with tempfile.TemporaryDirectory() as directory:
            provider = SequenceProvider()
            pipeline = ExamGenerationPipeline(provider, PipelineConfig(
                state_dir=Path(directory), max_repairs=0, question_workers=1,
                provider_attempts=2, provider_timeout_seconds=60,
            ))
            result = pipeline.generate("Italic source lead. Centered bold heading.", {"exam_type": "Reading"})

        self.assertEqual(result.status, "READY_FOR_REVIEW")
        self.assertEqual(result.validation_state, "PASS")
        self.assertEqual(result.round_trip_state, "PASS")
        self.assertEqual(provider.generate_calls, 2)

    def test_terminal_malformed_json_keeps_private_response_diagnostic(self):
        class AlwaysMalformedProvider(TextProvider):
            def __init__(self):
                self.generate_calls = 0

            def complete(self, **kwargs):
                if kwargs.get("phase") == "structure_outline":
                    self.generate_calls += 1
                    return '{"title":'
                raise AssertionError(kwargs.get("phase"))

        with tempfile.TemporaryDirectory() as directory:
            provider = AlwaysMalformedProvider()
            pipeline = ExamGenerationPipeline(provider, PipelineConfig(
                state_dir=Path(directory), max_repairs=0, question_workers=0,
                provider_attempts=2, provider_timeout_seconds=60,
            ))
            result = pipeline.generate("A short reading source.", {"exam_type": "Reading"})

        self.assertEqual(provider.generate_calls, 2)
        self.assertEqual(result.status, "FAILED")
        self.assertEqual(result.round_trip_state, "NOT_RUN")
        self.assertEqual(result.provider_response_debug, '{"title":')
        self.assertFalse(result.output_path)

    def test_pipeline_does_not_retry_strict_canonical_schema_failure(self):
        class InvalidSchemaProvider(TextProvider):
            def __init__(self):
                self.generate_calls = 0
                self.raw_response = ""

            def complete(self, **kwargs):
                if kwargs.get("phase") == "structure_outline":
                    return json.dumps(source_outline(kwargs["user"]))
                if kwargs.get("phase") == "structure_group":
                    self.generate_calls += 1
                    payload = structure_payload()
                    payload["unexpected"] = "must fail closed"
                    self.raw_response = json.dumps(payload)
                    return self.raw_response
                raise AssertionError(kwargs.get("phase"))

        with tempfile.TemporaryDirectory() as directory:
            provider = InvalidSchemaProvider()
            pipeline = ExamGenerationPipeline(provider, PipelineConfig(
                state_dir=Path(directory), max_repairs=0, question_workers=0,
                provider_attempts=2, provider_timeout_seconds=60,
            ))
            result = pipeline.generate("Italic source lead. Centered bold heading.", {"exam_type": "Reading"})

        self.assertEqual(provider.generate_calls, 1)
        self.assertEqual(result.status, "MANUAL_REVIEW")
        self.assertEqual(result.validation_state, "FAIL")
        self.assertEqual(result.round_trip_state, "NOT_RUN")
        self.assertFalse(result.output_path)
        self.assertEqual(result.issues[0]["code"], "STRUCTURE_SCHEMA_INVALID")
        self.assertEqual(result.provider_response_debug, provider.raw_response)

    def test_writing_numeric_schema_error_is_reported_as_generation_failure(self):
        class InvalidWritingProvider(TextProvider):
            def complete(self, **kwargs):
                if kwargs.get("phase") == "structure_group":
                    scope = json.loads(kwargs["user"].split("<SCOPE>", 1)[1].split("</SCOPE>", 1)[0])
                    task_number = scope["task_number"]
                    return json.dumps({
                        "title": "Writing source",
                        "exam_type": "Writing",
                        "time_limit": 60,
                        "sections": [],
                        "writing_tasks": [{"task_number": task_number, "prompt": [f"Prompt {task_number}"],
                            "minimum_words": "not-a-number" if task_number == 1 else 250,
                            "recommended_minutes": 20 if task_number == 1 else 40}],
                    })
                raise AssertionError(kwargs.get("phase"))

        with tempfile.TemporaryDirectory() as directory:
            pipeline = ExamGenerationPipeline(InvalidWritingProvider(), PipelineConfig(
                state_dir=Path(directory), max_repairs=0, question_workers=0,
                provider_attempts=1, provider_timeout_seconds=60,
            ))
            result = pipeline.generate(
                "WRITING TASK 1\nPrompt 1\nWRITING TASK 2\nPrompt 2",
                {"exam_type": "Writing"},
            )

        self.assertEqual(result.status, "MANUAL_REVIEW")
        self.assertEqual(result.issues[0]["code"], "STRUCTURE_SCHEMA_INVALID")
        self.assertIn("minimum_words", result.issues[0]["message"])

    def test_solver_contract_failure_keeps_private_question_response(self):
        class IncompleteSolverProvider(TextProvider):
            def __init__(self):
                self.incomplete_response = json.dumps({"question_number": 1, "correct_answers": ["first answer"]})

            def complete(self, **kwargs):
                if kwargs.get("phase") == "structure_outline":
                    return json.dumps(source_outline(kwargs["user"]))
                if kwargs.get("phase") == "structure_group":
                    return json.dumps(structure_payload())
                if kwargs.get("phase") == "question":
                    return self.incomplete_response
                raise AssertionError(kwargs.get("phase"))

        with tempfile.TemporaryDirectory() as directory:
            provider = IncompleteSolverProvider()
            pipeline = ExamGenerationPipeline(provider, PipelineConfig(
                state_dir=Path(directory), max_repairs=0, question_workers=1,
                provider_attempts=1, provider_timeout_seconds=60,
            ))
            result = pipeline.generate("Italic source lead. Centered bold heading.", {"exam_type": "Reading"})

        self.assertEqual(result.status, "MANUAL_REVIEW")
        self.assertEqual(result.issues[0]["code"], "EXPLANATION_FAILED")
        self.assertEqual(result.question_response_debug["solve:1"], provider.incomplete_response)

    def test_structure_generation_defers_answers_to_the_grounded_solver_stage(self):
        class StructureThenSolverProvider(TextProvider):
            def __init__(self):
                self.structure_system = ""

            def complete(self, **kwargs):
                phase = kwargs.get("phase")
                if phase == "structure_outline":
                    return json.dumps(source_outline(kwargs["user"]))
                if phase == "structure_group":
                    self.structure_system = str(kwargs.get("system") or "")
                    return json.dumps(structure_payload())
                if phase == "question":
                    question = json.loads(kwargs["user"].split("<QUESTION>", 1)[1].split("</QUESTION>", 1)[0])
                    source_question = reading_payload()["sections"][0]["questions"][question["question_number"] - 1]
                    return json.dumps({
                        "question_number": question["question_number"],
                        "correct_answers": source_question["correct_answers"],
                        "explanation": source_question["explanation"],
                    })
                if phase == "critic":
                    return '{"issues": [{"severity":"PASS","code":"CRITIC_PASS","message":"Verified."}]}'
                raise AssertionError(phase)

        with tempfile.TemporaryDirectory() as directory:
            provider = StructureThenSolverProvider()
            pipeline = ExamGenerationPipeline(provider, PipelineConfig(
                state_dir=Path(directory), max_repairs=0, question_workers=1,
                provider_attempts=1, provider_timeout_seconds=60,
            ))
            result = pipeline.generate("Italic source lead. Centered bold heading.", {"exam_type": "Reading"})

        self.assertIn("Do not solve or explain", provider.structure_system)
        self.assertIn("correct_answers to []", GENERATOR_SYSTEM)
        self.assertIn("must omit explanation", GENERATOR_SYSTEM)
        self.assertEqual(result.status, "READY_FOR_REVIEW")
        self.assertEqual(result.solved_count, 2)
        self.assertEqual(result.explanation_count, 2)

    def test_transient_question_failure_resumes_only_missing_question_from_checkpoint(self):
        class ResumeProvider(CompletePipelineProvider):
            def __init__(self):
                super().__init__()
                self.question_calls = {1: 0, 2: 0}

            def complete(self, **kwargs):
                if kwargs.get("phase") == "question":
                    question = question_from_call(kwargs)
                    number = question["question_number"]
                    self.question_calls[number] += 1
                    if number == 1 and self.question_calls[number] == 1:
                        raise ProviderError("Provider request timed out.", code="PROVIDER_TIMEOUT")
                return super().complete(**kwargs)

        with tempfile.TemporaryDirectory() as directory:
            provider = ResumeProvider()
            pipeline = ExamGenerationPipeline(provider, PipelineConfig(
                state_dir=Path(directory), max_repairs=1, question_workers=2,
                provider_attempts=1, provider_timeout_seconds=60,
            ))
            checkpoint = {}

            def save(value):
                checkpoint.clear()
                checkpoint.update(value)

            first = pipeline.generate("Italic source lead. Centered bold heading.",
                                      {"exam_type": "Reading"}, use_cache=False, on_checkpoint=save)
            second = pipeline.generate("Italic source lead. Centered bold heading.",
                                       {"exam_type": "Reading"}, use_cache=False,
                                       checkpoint=checkpoint, on_checkpoint=save)

        self.assertEqual(first.status, "FAILED")
        self.assertEqual([item["code"] for item in first.issues], ["PROVIDER_TIMEOUT"])
        self.assertEqual(first.solved_count, 1)
        self.assertEqual(second.status, "READY_FOR_REVIEW")
        self.assertEqual(provider.question_calls, {1: 2, 2: 1})
        self.assertIn("structure", checkpoint)

    def test_answer_is_preserved_while_only_explanation_is_repaired(self):
        class ExplanationRepairProvider(CompletePipelineProvider):
            def __init__(self):
                super().__init__()
                self.explanation_repairs = []

            def complete(self, **kwargs):
                if kwargs.get("phase") == "question":
                    question = question_from_call(kwargs)
                    expected = reading_payload()["sections"][0]["questions"][question["question_number"] - 1]
                    if question["question_number"] == 1:
                        return json.dumps({"question_number": 1, "correct_answers": expected["correct_answers"]})
                if kwargs.get("phase") == "explanation_repair":
                    question = question_from_call(kwargs)
                    self.explanation_repairs.append(question["question_number"])
                return super().complete(**kwargs)

        with tempfile.TemporaryDirectory() as directory:
            provider = ExplanationRepairProvider()
            result = ExamGenerationPipeline(provider, PipelineConfig(
                state_dir=Path(directory), max_repairs=1, question_workers=2,
                provider_attempts=1, provider_timeout_seconds=60,
            )).generate("Italic source lead. Centered bold heading.", {"exam_type": "Reading"}, use_cache=False)

        self.assertEqual(result.status, "READY_FOR_REVIEW")
        self.assertEqual(provider.explanation_repairs, [1])
        self.assertEqual(result.exam["sections"][0]["questions"][0]["correct_answers"], ["second answer"])

    def test_generate_batch_isolates_unexpected_file_failure_and_preserves_order(self):
        class IsolatedPipeline(ExamGenerationPipeline):
            def generate(self, source, requirements=None, **kwargs):
                if source == "broken-source":
                    raise RuntimeError("fixture crash")
                return PipelineResult(job_id=source, status="READY_FOR_REVIEW", original_filename=source)

        with tempfile.TemporaryDirectory() as directory:
            pipeline = IsolatedPipeline(CompletePipelineProvider(), PipelineConfig(state_dir=Path(directory), max_workers=2))
            results = pipeline.generate_batch(["first", "broken-source", "third"])

        self.assertEqual([item.original_filename for item in results], ["first", "broken-source", "third"])
        self.assertEqual([item.status for item in results], ["READY_FOR_REVIEW", "FAILED", "READY_FOR_REVIEW"])
        self.assertEqual(results[1].issues[0]["code"], "WORKER_FAILURE")

    def test_raw_docx_batch_keeps_order_and_uses_bounded_question_workers(self):
        class ConcurrentProvider(TextProvider):
            def __init__(self):
                self.lock = threading.Lock()
                self.active = 0
                self.peak = 0

            def complete(self, **kwargs):
                phase = kwargs.get("phase")
                if phase == "structure_outline":
                    return json.dumps(source_outline(kwargs["user"]))
                if phase == "structure_group":
                    return json.dumps(structure_payload())
                if phase == "critic":
                    return '{"issues": [{"severity":"PASS","code":"CRITIC_PASS","message":"Verified."}]}'
                if phase == "question":
                    question = json.loads(kwargs["user"].split("<QUESTION>", 1)[1].split("</QUESTION>", 1)[0])
                    number = question["question_number"]
                    with self.lock:
                        self.active += 1
                        self.peak = max(self.peak, self.active)
                    time.sleep(0.03)
                    with self.lock:
                        self.active -= 1
                    payload = reading_payload()["sections"][0]["questions"][number - 1]
                    return json.dumps({
                        "question_number": number,
                        "correct_answers": payload["correct_answers"],
                        "explanation": payload["explanation"],
                    })
                raise AssertionError(phase)

        with tempfile.TemporaryDirectory() as directory:
            paths = []
            for name in ("first.docx", "second.docx"):
                document = docx.Document()
                document.add_paragraph("[TYPE] Reading")
                document.add_paragraph("Italic source lead. Centered bold heading.")
                path = Path(directory) / name
                document.save(path)
                paths.append(str(path))
            provider = ConcurrentProvider()
            pipeline = ExamGenerationPipeline(provider, PipelineConfig(
                state_dir=Path(directory) / "state", max_workers=2, question_workers=2, max_repairs=0,
            ))
            results = pipeline.transform_raw_docx_batch(paths)

            self.assertEqual([result.original_filename for result in results], ["first.docx", "second.docx"])
            self.assertTrue(all(result.status == "READY_FOR_REVIEW" for result in results))
            self.assertTrue(all(result.question_count == 2 for result in results))
            self.assertTrue(all(result.solved_count == 2 for result in results))
            self.assertLessEqual(provider.peak, 4)  # 2 files × at most 2 question workers each
            self.assertGreaterEqual(provider.peak, 2)


    def test_embedded_media_is_bound_reembedded_and_parser_verified(self):
        tiny_png = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/ScL9ZQAAAABJRU5ErkJggg=="
        )
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            image = root / "chart.png"
            image.write_bytes(tiny_png)
            raw = root / "raw.docx"
            source = docx.Document()
            source.add_paragraph("[TYPE] Writing")
            source.add_picture(str(image))
            source.save(raw)
            media_store = InternalMediaStore(root / "media")
            extraction = extract_docx_media(raw, "raw-source", media_store)
            self.assertEqual(len(extraction.occurrences), 1)
            occurrence = extraction.occurrences[0]
            self.assertEqual(media_store.get(occurrence.asset_id).asset_id, occurrence.asset_id)
            self.assertNotIn("storage_path", extraction.to_dict()["assets"][0])
            payload = {
                "title": "Writing chart with managed media",
                "exam_type": "Writing",
                "time_limit": 60,
                "sections": [],
                "writing_tasks": [
                    {"task_number": 1, "prompt": [
                        "Describe the chart.",
                        {"media": {"occurrence_id": occurrence.occurrence_id, "asset_id": occurrence.asset_id}},
                        "Compare its main features.",
                    ], "minimum_words": 150, "recommended_minutes": 20,
                     "media_url": "https://evil.example/tracker.png", "media_asset_id": occurrence.asset_id},
                    {"task_number": 2, "prompt": ["Discuss the topic."], "minimum_words": 250, "recommended_minutes": 40},
                ],
                "media_bindings": [{
                    "occurrence_id": occurrence.occurrence_id,
                    "asset_id": occurrence.asset_id,
                    "association": "writing_task:1",
                    "required_for_solving": True,
                }],
            }
            exam = CanonicalExam.from_dict(payload)
            paths, diagnostics = resolve_media_bindings(exam, extraction, "https://ielts-os.example")
            self.assertEqual(diagnostics, [])
            self.assertEqual(exam.writing_tasks[0].media_url,
                             f"https://ielts-os.example/api/exam-media/{occurrence.asset_id}")
            output = render_docx(exam, root / "formatted.docx", media_asset_paths=paths)
            verify_embedded_media_roundtrip(output, paths)
            verify_docx_roundtrip(exam, output)

    def test_visual_question_solver_and_critic_receive_only_bound_image(self):
        tiny_png = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/ScL9ZQAAAABJRU5ErkJggg=="
        )

        class VisionProvider(TextProvider):
            def __init__(self):
                self.calls = []

            def complete(self, **kwargs):
                self.calls.append(kwargs)
                if kwargs.get("phase") == "critic":
                    return json.dumps({"issues": [{"severity": "PASS", "code": "VISUAL_PASS",
                                                   "message": "The chart evidence is visible.", "question_number": 1}]})
                return json.dumps({
                    "question_number": 1,
                    "correct_answers": ["second answer"],
                    "explanation": {
                        "why": "The chart visibly marks the second answer at the requested location.",
                        "evidence": {"quote": "second answer", "source": "image",
                                     "source_id": "media_occ_1", "locator": "upper-right label"},
                        "paraphrase_mapping": "upper-right label -> second answer",
                        "correct_reason": "The requested chart position carries the second label.",
                        "distractor_reasons": {"first answer": "It is shown at a different chart position."},
                        "student_takeaway": "Trace the requested position before matching its label.",
                    },
                })

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            image = root / "chart.png"
            image.write_bytes(tiny_png)
            asset = MediaAsset("asset_1", "abc123", "chart.png", "image/png", "png",
                               len(tiny_png), 1, 1, str(image))
            occurrence = MediaOccurrence("media_occ_1", "asset_1", "source_1", "word/document.xml",
                                         "rId1", 1, "body/1", None, None, None, {}, "chart", "", "")
            extraction = MediaExtraction([asset], [occurrence], [])
            payload = reading_payload()
            payload["sections"][0]["questions"] = payload["sections"][0]["questions"][:1]
            question_payload = payload["sections"][0]["questions"][0]
            question_payload["correct_answers"] = []
            question_payload.pop("explanation", None)
            payload["media_bindings"] = [{
                "occurrence_id": "media_occ_1", "asset_id": "asset_1",
                "association": "question:1", "required_for_solving": True,
            }]
            exam = CanonicalExam.from_dict(payload)
            provider = VisionProvider()
            pipeline = ExamGenerationPipeline(provider, PipelineConfig(
                state_dir=root / "state", max_repairs=0, question_workers=1,
            ))

            issues = pipeline._solve_questions(exam, "ordinary source text", None, extraction=extraction)

            self.assertEqual(issues, [])
            solver_call = next(call for call in provider.calls if call.get("phase") == "question")
            self.assertEqual(len(solver_call["image_inputs"]), 1)
            self.assertEqual(solver_call["image_inputs"][0]["occurrence_id"], "media_occ_1")
            self.assertNotIn(str(image), solver_call["user"])
            self.assertEqual(ground_exam(exam, "ordinary source text", extraction), [])
            self.assertEqual(exam.sections[0].questions[0].explanation.evidence.source_hash, "abc123")

            review = critic_review(provider, exam, "ordinary source text", "critic-model",
                                   image_inputs=solver_call["image_inputs"], focus_question_number=1)
            self.assertEqual(review[0].severity, "PASS")
            critic_call = next(call for call in provider.calls if call.get("phase") == "critic")
            self.assertEqual(critic_call["image_inputs"][0]["occurrence_id"], "media_occ_1")

    def test_media_binding_required_for_solving_is_strict_boolean(self):
        payload = reading_payload()
        payload["media_bindings"] = [{
            "occurrence_id": "media_occ_1", "asset_id": "asset_1",
            "association": "question:1", "required_for_solving": "false",
        }]
        with self.assertRaisesRegex(CanonicalSchemaError, "boolean"):
            CanonicalExam.from_dict(payload)


if __name__ == "__main__":
    unittest.main()
