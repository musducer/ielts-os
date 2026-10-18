import unittest
import docx
from api.index import process_block


class TestAllowOptionReuse(unittest.TestCase):
    def test_reuse_tag_enables_allow_option_reuse(self):
        doc = docx.Document()
        lines = [
            doc.add_paragraph("[DRAG]"),
            doc.add_paragraph("[REUSE_OPTIONS]"),
            doc.add_paragraph("Which category describes each DVD?"),
            doc.add_paragraph("OPTIONS:"),
            doc.add_paragraph("A. Current"),
            doc.add_paragraph("B. Out of date"),
            doc.add_paragraph("C. Planned"),
            doc.add_paragraph("21. DVD One"),
            doc.add_paragraph("*A"),
            doc.add_paragraph("22. DVD Two"),
            doc.add_paragraph("*A"),
        ]

        target_questions = []
        process_block("DRAG", lines, target_questions)
        self.assertEqual(len(target_questions), 2)
        for q in target_questions:
            self.assertTrue(q.get("allowOptionReuse"), f"Expected allowOptionReuse to be True on {q}")

    def test_instruction_phrase_enables_allow_option_reuse(self):
        doc = docx.Document()
        lines = [
            doc.add_paragraph("[DRAG]"),
            doc.add_paragraph("You may use any letter more than once."),
            doc.add_paragraph("OPTIONS:"),
            doc.add_paragraph("A. Current"),
            doc.add_paragraph("B. Out of date"),
            doc.add_paragraph("21. DVD One"),
            doc.add_paragraph("*A"),
        ]

        target_questions = []
        process_block("DRAG", lines, target_questions)
        self.assertEqual(len(target_questions), 1)
        self.assertTrue(target_questions[0].get("allowOptionReuse"))

    def test_duplicate_answers_enable_allow_option_reuse(self):
        doc = docx.Document()
        lines = [
            doc.add_paragraph("[MATCHING]"),
            doc.add_paragraph("Match the DVD to its category."),
            doc.add_paragraph("OPTIONS:"),
            doc.add_paragraph("A. Current"),
            doc.add_paragraph("B. Out of date"),
            doc.add_paragraph("C. Planned"),
            doc.add_paragraph("21. DVD One"),
            doc.add_paragraph("*Current"),
            doc.add_paragraph("22. DVD Two"),
            doc.add_paragraph("*Current"),
        ]

        target_questions = []
        process_block("MATCHING", lines, target_questions)
        self.assertEqual(len(target_questions), 2)
        for q in target_questions:
            self.assertTrue(q.get("allowOptionReuse"))

    def test_flow_drag_title_and_steps(self):
        doc = docx.Document()
        p_title = doc.add_paragraph()
        run = p_title.add_run("Plan for last lesson")
        run.bold = True
        lines = [
            doc.add_paragraph("[FLOW_DRAG]"),
            doc.add_paragraph("Questions 26-30"),
            doc.add_paragraph("Complete the flow-chart below."),
            p_title,
            doc.add_paragraph("↓"),
            doc.add_paragraph("(class) watch a movie"),
            doc.add_paragraph("↓"),
            doc.add_paragraph("(small groups) have a discussion"),
            doc.add_paragraph("↓"),
            doc.add_paragraph("(class) [26]"),
            doc.add_paragraph("↓"),
            doc.add_paragraph("(pairs) [27]"),
            doc.add_paragraph("OPTIONS:"),
            doc.add_paragraph("A. ask questions"),
            doc.add_paragraph("B. spell key words"),
            doc.add_paragraph("26. "),
            doc.add_paragraph("*ask questions"),
            doc.add_paragraph("27. "),
            doc.add_paragraph("*spell key words"),
        ]

        target_questions = []
        process_block("FLOW_DRAG", lines, target_questions)
        self.assertEqual(len(target_questions), 2)
        self.assertIn("Plan for last lesson", target_questions[0].get("flowTitle", ""))
        steps = target_questions[0].get("flowSteps", [])
        self.assertEqual(len(steps), 4)
        self.assertIsNone(steps[0]["questionNumber"])
        self.assertIn("watch a movie", steps[0]["html"])
        self.assertIsNone(steps[1]["questionNumber"])
        self.assertIn("have a discussion", steps[1]["html"])
        self.assertEqual(steps[2]["questionNumber"], 26)
        self.assertEqual(steps[3]["questionNumber"], 27)


if __name__ == '__main__':
    unittest.main()
