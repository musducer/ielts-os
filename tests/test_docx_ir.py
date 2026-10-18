import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from PIL import Image

from exam_generator.docx_ir import IR_VERSION
from exam_generator.media import InternalMediaStore, extract_docx_media
from exam_generator.source import load_source_snapshot


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
V = "urn:schemas-microsoft-com:vml"


def _leaf_sequence(snapshot):
    return ["IMAGE" if block["kind"] == "IMAGE" else block["visible_text"]
            for block in snapshot.blocks
            if not block.get("children") and block["coverage_state"] != "IGNORED_DECORATION"
            and (block["visible_text"] or block["kind"] == "IMAGE")]


class OrderedDocxIRTests(unittest.TestCase):
    def test_true_order_nested_table_and_media_binding_are_source_addressed(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            picture = root / "pixel.png"
            Image.new("RGB", (8, 5), "navy").save(picture)
            document = Document()
            document.add_paragraph("Before image")
            document.add_paragraph().add_run().add_picture(str(picture))
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "Outer cell"
            nested = table.cell(0, 0).add_table(rows=1, cols=1)
            nested.cell(0, 0).text = "Nested cell"
            document.add_paragraph("After table")
            source = root / "ordered.docx"
            document.save(source)

            snapshot = load_source_snapshot(str(source))
            self.assertEqual(snapshot.ir_version, IR_VERSION)
            self.assertEqual(snapshot.coverage["state"], "PASS")
            self.assertEqual(
                _leaf_sequence(snapshot),
                ["Before image", "IMAGE", "Outer cell", "Nested cell", "After table"],
            )
            image = next(block for block in snapshot.blocks if block["kind"] == "IMAGE")
            self.assertEqual(image["nearby_text"], {"before": "Before image", "after": "Outer cell"})
            self.assertEqual((image["width_emu"] is not None, image["height_emu"] is not None), (True, True))
            self.assertEqual(image["drawing_type"], "inline")

            extraction = extract_docx_media(source, snapshot.source_hash,
                                            InternalMediaStore(root / "media"))
            snapshot.bind_media(extraction)
            image_inputs = snapshot.provider_image_inputs([image["occurrence_id"]])
            self.assertEqual(image_inputs[0]["occurrence_id"], image["occurrence_id"])
            self.assertTrue(image_inputs[0]["data"])
            self.assertNotIn("path", snapshot.prompt_payload())

            copied = root / "renamed.docx"
            copied.write_bytes(source.read_bytes())
            second = load_source_snapshot(str(copied))
            self.assertEqual(snapshot.source_hash, second.source_hash)
            self.assertEqual([block["block_id"] for block in snapshot.blocks],
                             [block["block_id"] for block in second.blocks])

    def test_inherited_styles_alternate_content_and_tracked_deletion(self):
        with tempfile.TemporaryDirectory() as directory:
            document = Document()
            base = document.styles.add_style("IR Base", WD_STYLE_TYPE.PARAGRAPH)
            base.font.bold = True
            base.font.italic = True
            derived = document.styles.add_style("IR Derived", WD_STYLE_TYPE.PARAGRAPH)
            derived.base_style = base
            derived.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            document.add_paragraph("Inherited format", style=derived)

            paragraph = document.add_paragraph()
            paragraph._p.append(parse_xml(
                f'<w:del xmlns:w="{W}"><w:r><w:delText>Deleted text</w:delText></w:r></w:del>'
            ))
            paragraph._p.append(parse_xml(
                f'<mc:AlternateContent xmlns:mc="{MC}" xmlns:w="{W}" xmlns:w14="{W14}">'
                '<mc:Choice Requires="w14"><w:r><w:t>Visible choice</w:t></w:r></mc:Choice>'
                '<mc:Fallback><w:r><w:t>Duplicate fallback</w:t></w:r></mc:Fallback>'
                '</mc:AlternateContent>'
            ))
            source = Path(directory) / "tracked.docx"
            document.save(source)

            snapshot = load_source_snapshot(str(source))
            inherited = next(block for block in snapshot.blocks
                             if block["visible_text"] == "Inherited format")
            self.assertEqual(inherited["alignment"], "center")
            self.assertTrue(inherited["runs"][0]["bold"])
            self.assertTrue(inherited["runs"][0]["italic"])
            self.assertIn("Visible choice", snapshot.visible_text)
            self.assertNotIn("Duplicate fallback", snapshot.visible_text)
            self.assertNotIn("Deleted text", snapshot.visible_text)
            deletion = next(block for block in snapshot.blocks
                            if block.get("ooxml_tag") == "del")
            self.assertEqual(deletion["coverage_state"], "IGNORED_DECORATION")

    def test_merged_and_nested_tables_preserve_geometry_without_manual_review(self):
        with tempfile.TemporaryDirectory() as directory:
            document = Document()
            table = document.add_table(rows=3, cols=3)
            table.cell(0, 0).merge(table.cell(0, 1)).text = "Wide heading"
            table.cell(1, 0).merge(table.cell(2, 0)).text = "Tall label"
            table.cell(1, 1).text = "Normal cell"
            table.cell(1, 1).add_table(rows=1, cols=1).cell(0, 0).text = "Nested detail"
            source = Path(directory) / "table.docx"
            document.save(source)

            snapshot = load_source_snapshot(str(source))
            self.assertEqual(snapshot.coverage["state"], "PASS")
            outer = next(block for block in snapshot.blocks if block["kind"] == "TABLE")
            self.assertEqual(outer["rows"][0]["cells"][0]["column_span"], 2)
            vertical_owner = outer["rows"][1]["cells"][0]
            self.assertEqual(vertical_owner["vertical_merge"], "restart")
            self.assertEqual(vertical_owner["row_span"], 2)
            continuation = outer["rows"][2]["cells"][0]
            self.assertEqual(continuation["vertical_merge"], "continue")
            self.assertIn("Nested detail", outer["visible_text"])
            self.assertGreaterEqual(sum(block["kind"] == "TABLE" for block in snapshot.blocks), 2)

    def test_textboxes_and_shape_text_are_normalized_but_unrenderable_chart_fails(self):
        with tempfile.TemporaryDirectory() as directory:
            document = Document()
            paragraph = document.add_paragraph()
            paragraph._p.append(parse_xml(
                f'<w:r xmlns:w="{W}" xmlns:v="{V}"><w:pict><v:shape><v:textbox>'
                '<w:txbxContent><w:p><w:r><w:t>Textbox content</w:t></w:r></w:p></w:txbxContent>'
                '</v:textbox></v:shape></w:pict></w:r>'
            ))
            shape = document.add_paragraph()
            shape._p.append(parse_xml(
                f'<w:r xmlns:w="{W}" xmlns:wp="{WP}" xmlns:a="{A}"><w:drawing><wp:inline>'
                '<a:graphic><a:graphicData><a:t>Shape label</a:t></a:graphicData></a:graphic>'
                '</wp:inline></w:drawing></w:r>'
            ))
            source = Path(directory) / "shapes.docx"
            document.save(source)
            snapshot = load_source_snapshot(str(source))
            self.assertEqual(snapshot.coverage["state"], "PASS")
            self.assertIn("Textbox content", snapshot.visible_text)
            self.assertIn("Shape label", snapshot.visible_text)
            self.assertTrue(any(block["kind"] == "TEXTBOX" for block in snapshot.blocks))

            chart_document = Document()
            chart = chart_document.add_paragraph()
            chart._p.append(parse_xml(
                f'<w:r xmlns:w="{W}" xmlns:wp="{WP}" xmlns:a="{A}" xmlns:c="{C}" xmlns:r="{R}">'
                '<w:drawing><wp:inline><a:graphic><a:graphicData><c:chart r:id="rIdMissing"/>'
                '</a:graphicData></a:graphic></wp:inline></w:drawing></w:r>'
            ))
            chart_source = Path(directory) / "chart.docx"
            chart_document.save(chart_source)
            chart_snapshot = load_source_snapshot(str(chart_source))
            self.assertEqual(chart_snapshot.coverage["state"], "FAIL")
            self.assertIn("SOURCE_DRAWING_CONTENT_UNSUPPORTED",
                          {item["code"] for item in chart_snapshot.diagnostics})


if __name__ == "__main__":
    unittest.main()
