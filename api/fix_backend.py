import os

def upgrade_backend():
    target_file = "index.py"
    if not os.path.exists(target_file):
        print("❌ LỖI: Không tìm thấy file index.py!")
        return

    new_backend_code = """import re
import io
import time
import traceback
from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
import docx
from docx.table import Table
from docx.text.paragraph import Paragraph
from typing import List, Dict, Any

app = FastAPI(docs_url="/api/docs", openapi_url="/api/openapi.json")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_credentials=True, allow_methods=["*"], allow_headers=["*"])

VALID_BLOCK_TYPES = ["BLANK", "CHOICE", "CHOICE_MULTIPLE", "MATCHING", "DRAG_DROP", "SHORT_ANSWER"]

RE_Q_START = re.compile(r'^(?:Question|Câu)\\s+\\d+', re.IGNORECASE)
RE_Q_NUMBER = re.compile(r'^\\d+[\\.\\)\\:]')
RE_TFNG = re.compile(r'\\btrue\\b[\\s\\S]*\\bfalse\\b[\\s\\S]*\\bnot\\s+given\\b', re.IGNORECASE)
RE_YNNG = re.compile(r'\\byes\\b[\\s\\S]*\\bno\\b[\\s\\S]*\\bnot\\s+given\\b', re.IGNORECASE)

def is_option(text: str) -> bool:
    text = text.strip()
    # Dạng 1 chuẩn: A. text, A) text
    if re.match(r'^[A-Za-z][\\.\\)]\\s+', text): return True
    # Dạng 2: A, B, C đứng độc lập
    if re.match(r'^[A-Za-z]$', text): return True
    # Dạng 3: A Text (List of people, features...)
    m = re.match(r'^([A-K])\\s+(.+)$', text)
    if m:
        letter, rest = m.groups()
        if len(text) < 80: # Giới hạn độ dài để tránh nhận nhầm câu văn bản dài
            if letter == 'A' and rest[0].islower(): return False # "A boy walked..."
            if letter == 'I' and (rest.lower().startswith('am ') or rest.lower().startswith('have ') or rest.lower().startswith('was ')): return False
            return True
    return False

def paragraph_to_html(para: Paragraph) -> str:
    if not para.text: return ""
    html_parts = []
    for run in para.runs:
        text = run.text or ""
        if not text: continue
        text = text.replace('\\t', '&nbsp;&nbsp;&nbsp;&nbsp;').replace('  ', '&nbsp;&nbsp;').replace('\\n', '<br/>')
        if run.bold: text = f"<strong>{text}</strong>"
        if run.italic: text = f"<em>{text}</em>"
        if run.underline: text = f"<u>{text}</u>"
        html_parts.append(text)
    p_html = "".join(html_parts)
    if not p_html.replace('&nbsp;', '').strip() and '<br/>' not in p_html:
        return ""
    
    align = para.alignment
    style_name = para.style.name.lower() if para.style else ""
    wrapper_style = ""
    if align == 1: wrapper_style += "text-align: center; "
    elif align == 2: wrapper_style += "text-align: right; "
    elif align == 3: wrapper_style += "text-align: justify; "
    
    # BẢO TỒN ĐỊNH DẠNG BULLET POINTS CHO NOTES COMPLETION
    is_list = False
    if 'list' in style_name or 'bullet' in style_name:
        is_list = True
    elif para._p.pPr is not None and para._p.pPr.numPr is not None:
        is_list = True

    if 'heading 1' in style_name: return f'<h1 style="{wrapper_style}margin:15px 0;font-size:24px;">{p_html}</h1>'
    if 'heading 2' in style_name: return f'<h2 style="{wrapper_style}margin:12px 0;font-size:20px;">{p_html}</h2>'
    if 'heading 3' in style_name: return f'<h3 style="{wrapper_style}margin:10px 0;font-size:18px;">{p_html}</h3>'
    
    if is_list:
        return f'<div style="display: list-item; margin-left: 24px; list-style-type: disc; {wrapper_style}">{p_html}</div>'

    return f'<div style="{wrapper_style}">{p_html}</div>' if wrapper_style else p_html

def table_to_html(table: Table) -> str:
    rows = table.rows
    if not rows: return ""
    n_rows, n_cols = len(rows), len(rows[0].cells)
    html = '<table style="width:100%; border-collapse: collapse; margin-bottom: 15px;"><tbody>'
    for r_idx in range(n_rows):
        html += "<tr>"
        for c_idx in range(n_cols):
            cell_text = "".join([f"<div>{paragraph_to_html(p)}</div>" for p in table.cell(r_idx, c_idx).paragraphs if p.text.strip()])
            html += f'<td style="border:1px solid #ccc; padding:8px;">{cell_text}</td>'
        html += "</tr>"
    html += "</tbody></table>"
    return html

def process_block(block_type: str, lines: List[Any], target_questions: List[Dict]):
    html_lines, text_lines = [], []
    for item in lines:
        if isinstance(item, Paragraph):
            h = paragraph_to_html(item)
            if not h: continue
            html_lines.append(h)
            text_lines.append(re.sub(r'<[^>]+>', '', h).strip())
        elif isinstance(item, Table):
            html_lines.append(table_to_html(item))
            text_lines.append("[TABLE]")
    
    questions_in_block = []
    current_q = None
    accumulated_instruction = ""
    group_context = ""
    pre_question_options = []
    in_context = False
    
    for h_line, t_line in zip(html_lines, text_lines):
        t_upper = t_line.upper()
        if "[CONTEXT]" in t_upper:
            in_context = True
            continue
        if "[/CONTEXT]" in t_upper:
            in_context = False
            continue
        
        if in_context:
            group_context += f"<div>{h_line}</div>"
            continue

        is_q_start = bool(RE_Q_START.match(t_line) or RE_Q_NUMBER.match(t_line))
        
        if not is_q_start and not current_q:
            if is_option(t_line):
                pre_question_options.append(t_line)
                continue
            accumulated_instruction += f"<div>{h_line}</div>"
            continue

        if is_q_start:
            if current_q: questions_in_block.append(current_q)
            content_only = re.sub(r'^(?:Question|Câu)?\\s*\\d+[\\.\\:\\)]?\\s*', '', t_line, flags=re.IGNORECASE).strip()
            current_q = {
                "type": block_type, "text": content_only if content_only else h_line,
                "options": pre_question_options.copy(), "correctAnswers": [],
                "instruction": accumulated_instruction,
                "groupContext": group_context
            }
            continue
        
        if current_q:
            if t_line.startswith("*"):
                current_q["correctAnswers"].append(re.sub(r'^\\*+\\s*', '', t_line).strip())
            elif is_option(t_line):
                current_q["options"].append(t_line)
            else:
                current_q["text"] += f" <div>{h_line}</div>"

    if current_q: questions_in_block.append(current_q)
    
    # Chia sẻ options cho các câu trong nhóm
    shared_options = []
    for q in questions_in_block:
        if q["options"]: shared_options = q["options"]
        elif shared_options and q["type"] in ["CHOICE", "CHOICE_MULTIPLE", "MATCHING"]:
            q["options"] = shared_options.copy()
    
    for q in questions_in_block:
        lower_txt = (q['instruction'] + q['text']).lower()
        sub_t = "SENTENCE"
        if "flow-chart" in lower_txt or "flowchart" in lower_txt: sub_t = "FLOWCHART"
        elif "summary" in lower_txt: sub_t = "SUMMARY"
        elif "note" in lower_txt: sub_t = "NOTES"
        
        is_tfng = RE_TFNG.search(lower_txt) or "true/false/not given" in lower_txt
        is_ynng = RE_YNNG.search(lower_txt) or "yes/no/not given" in lower_txt
        
        if not q["options"]:
            if is_tfng: q["options"] = ["TRUE", "FALSE", "NOT GIVEN"]
            elif is_ynng: q["options"] = ["YES", "NO", "NOT GIVEN"]
            
        output_type = "DRAG_DROP_HEADING" if "heading" in lower_txt and q["type"] == "MATCHING" else q["type"]
        if is_tfng or is_ynng: output_type = "CHOICE"
        
        final_correct_answer = ""
        # THUẬT TOÁN ÁNH XẠ ĐÁP ÁN (SMART MAPPER)
        if output_type in ["CHOICE", "MATCHING"]:
            if q["correctAnswers"] and q["options"]:
                ans_str = q["correctAnswers"][0].strip().upper()
                for idx, opt in enumerate(q["options"]):
                    opt_str = opt.strip()
                    m_letter = re.match(r'^([A-Za-z])(?:[\\.\\)]\\s+|\\s+|$)', opt_str)
                    opt_letter = m_letter.group(1).upper() if m_letter else chr(65 + idx)
                    
                    if ans_str == opt_letter or ans_str == opt_str.upper():
                        final_correct_answer = idx
                        break
                if final_correct_answer == "": final_correct_answer = 0
            else:
                final_correct_answer = 0
        elif output_type == "CHOICE_MULTIPLE":
            ans_arr = []
            if q["correctAnswers"]:
                for ans in q["correctAnswers"]:
                    ans_str = ans.strip().upper()
                    for idx, opt in enumerate(q["options"] if q["options"] else []):
                        opt_str = opt.strip()
                        m_letter = re.match(r'^([A-Za-z])(?:[\\.\\)]\\s+|\\s+|$)', opt_str)
                        opt_letter = m_letter.group(1).upper() if m_letter else chr(65 + idx)
                        if ans_str == opt_letter or ans_str == opt_str.upper():
                            ans_arr.append(idx)
                            break
            final_correct_answer = ans_arr
        else:
            final_correct_answer = q["correctAnswers"][0] if q["correctAnswers"] else ""

        target_questions.append({
            "id": f"q_{int(time.time() * 1000)}_{len(target_questions)}",
            "type": output_type,
            "subType": sub_t, 
            "instruction": q["instruction"], 
            "groupContext": q["groupContext"],
            "text": q["text"], 
            "options": q["options"] if q["options"] else [],
            "correctAnswer": final_correct_answer
        })

def parse_docx_to_quiz(doc):
    title, time_limit, quiz_type, audio_url = "Untitled Mock Test", 60, "Reading", ""
    sections = []
    
    current_passage = ""
    current_questions = []
    
    state = "META"
    block_type = None
    current_block_lines = []

    def flush_block():
        nonlocal current_block_lines, block_type
        if block_type and current_block_lines:
            process_block(block_type, current_block_lines, current_questions)
        current_block_lines = []

    for element in doc.element.body:
        if element.tag.endswith('p'):
            para = Paragraph(element, doc)
            text = para.text.strip()
            if not text: continue
            text_upper = text.upper()
            
            if "[TITLE]" in text_upper: title = re.sub(r'\\[TITLE\\]', '', text, flags=re.IGNORECASE).strip(); continue
            if "[TIME]" in text_upper: 
                try: time_limit = int(re.sub(r'\\D', '', text))
                except: pass
                continue
            if "[TYPE]" in text_upper: quiz_type = re.sub(r'\\[TYPE\\]', '', text, flags=re.IGNORECASE).strip(); continue
            if "[AUDIO]" in text_upper: audio_url = re.sub(r'\\[AUDIO\\]', '', text, flags=re.IGNORECASE).strip(); continue
            
            if "[PASSAGE]" in text_upper:
                flush_block()
                if current_passage or current_questions:
                    sections.append({"passage": current_passage, "questions": current_questions})
                current_passage, current_questions = "", []
                state = "PASSAGE"
                block_type = None
                continue
            
            if "[QUESTIONS]" in text_upper:
                flush_block()
                state = "QUESTIONS"
                block_type = None
                continue

            if state == "PASSAGE":
                html = paragraph_to_html(para)
                if html: current_passage += f"<div style='margin-bottom:12px;'>{html}</div>"
            elif state == "QUESTIONS":
                match = re.search(r'\\[(/?\\w+)\\]', text_upper)
                if match and match.group(1) in VALID_BLOCK_TYPES:
                    flush_block()
                    block_type = match.group(1)
                    continue
                
                if block_type:
                    current_block_lines.append(para)
                
        elif element.tag.endswith('tbl'):
            if state == "PASSAGE":
                current_passage += table_to_html(Table(element, doc))
            elif state == "QUESTIONS" and block_type:
                current_block_lines.append(Table(element, doc))

    flush_block()
    if current_passage or current_questions:
        sections.append({"passage": current_passage, "questions": current_questions})

    all_questions = []
    for sec in sections: all_questions.extend(sec["questions"])

    return {
        "id": f"quiz_{int(time.time() * 1000)}",
        "title": title,
        "type": quiz_type,
        "audioUrl": audio_url,
        "timeLimit": time_limit,
        "passage": sections[0]["passage"] if sections else "",
        "sections": sections,
        "questions": all_questions,
        "active": False,
        "maxAttempts": 1,
        "audience": "ALL",
        "targetStudentIds": [],
        "folder": "Root",
        "isSEBRequired": False,
        "passcode": "",
        "scheduledStart": "",
        "scheduledEnd": ""
    }

@app.post("/api/upload_docx")
async def upload_docx(file: UploadFile = File(...)):
    try:
        doc = docx.Document(io.BytesIO(await file.read()))
        quiz = parse_docx_to_quiz(doc)
        if not quiz.get("questions") or len(quiz["questions"]) == 0:
            return {"success": False, "error": "Lỗi: Không trích xuất được câu hỏi nào. Hãy chắc chắn bạn đã gắn đủ thẻ [PASSAGE] và [QUESTIONS]."}
        return {"success": True, "quiz": quiz}
    except Exception as e:
        return {"success": False, "error": f"Lỗi xử lý file: {str(e)}\\n{traceback.format_exc()}"}

@app.get("/api/health")
async def health_check():
    return {"status": "ok", "timestamp": time.time()}
"""
    with open(target_file, "w", encoding="utf-8") as f:
        f.write(new_backend_code)
    
    print("✅ ĐÃ FIX BACKEND THÀNH CÔNG! (Giải quyết Bullet Points, List Options và Smart Answer Mapper)")
    print("⚠️ BẠN HÃY KHỞI ĐỘNG LẠI SERVER FASTAPI ĐỂ NHẬN CODE MỚI NHÉ!")

if __name__ == "__main__":
    upgrade_backend()