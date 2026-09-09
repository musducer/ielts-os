"""Hard parser-contract gate for generated DOCX.

This module calls ``api.index.parse_docx_to_quiz`` directly.  A document that
fails this comparison is not eligible for import or publication.
"""

from __future__ import annotations

import html
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List

import docx

from .schema import CanonicalExam, Option, Question, RichParagraph, option_answer_values


class ParserContractError(ValueError):
    pass


def _plain(value: Any) -> str:
    if isinstance(value, RichParagraph):
        value = value.text
    elif isinstance(value, Option):
        value = value.content.text
    elif isinstance(value, dict):
        value = RichParagraph.from_value(value).text
    value = html.unescape(re.sub(r"<[^>]+>", " ", str(value or "")))
    return re.sub(r"\s+", " ", value).strip()


def _clean_option(value: Any) -> str:
    return re.sub(r"^[A-Za-z][.)]\s*", "", _plain(value)).strip()


def _canonical_choice_answer(question: Question) -> Any:
    raw = option_answer_values(question)
    indexes: List[int] = []
    for answer in raw:
        answer = _plain(answer)
        if re.fullmatch(r"[A-Za-z]", answer):
            indexes.append(ord(answer.upper()) - 65)
            continue
        for index, option in enumerate(question.options):
            if _clean_option(option).casefold() == answer.casefold():
                indexes.append(index)
                break
    return indexes if question.block_type == "CHOICE_MULTIPLE" else (indexes[0] if indexes else None)


def _canonical_answer(question: Question) -> Any:
    if question.block_type in {"CHOICE", "CHOICE_MULTIPLE"}:
        return _canonical_choice_answer(question)
    answer = _plain(option_answer_values(question)[0] if option_answer_values(question) else "")
    if question.block_type in {"MATCHING", "DRAG", "DRAG_DROP", "MAP_DRAG", "FLOW_DRAG"}:
        return _clean_option(answer)
    return answer


def _parser_type(question: Question) -> str:
    if question.block_type == "SHORT_ANSWER":
        return "BLANK"
    if question.block_type == "DRAG":
        return "DRAG_DROP"
    return question.block_type


def _canonical_explanation_raw(question: Question) -> str:
    if not question.explanation:
        return ""
    chunks = [_plain(question.explanation.why)]
    evidence = question.explanation.evidence
    if evidence.quote:
        chunks.append(f'"{_plain(evidence.quote)}"')
    if evidence.locator:
        locator = _plain(evidence.locator)
        chunks.append(locator if locator.startswith("[") else f"[{locator}]")
    return " ".join(chunk for chunk in chunks if chunk)


def canonical_parser_projection(exam: CanonicalExam) -> Dict[str, Any]:
    if exam.exam_type == "Writing":
        return {
            "title": _plain(exam.title),
            "type": "Writing",
            "timeLimit": exam.time_limit,
            "writingTasks": [
                {
                    "taskNumber": task.task_number,
                    "title": _plain(task.title) or f"Writing Task {task.task_number}",
                    "instructions": _plain(task.instructions) or f"You should spend about {task.recommended_minutes or (20 if task.task_number == 1 else 40)} minutes on this task. Write at least {task.minimum_words or (150 if task.task_number == 1 else 250)} words.",
                    "minimumWords": task.minimum_words or (150 if task.task_number == 1 else 250),
                    "recommendedMinutes": task.recommended_minutes or (20 if task.task_number == 1 else 40),
                    "mediaUrl": task.media_url,
                    "prompt": " ".join(_plain(line) for line in task.prompt),
                }
                for task in sorted(exam.writing_tasks, key=lambda item: item.task_number)
            ],
        }
    sections = []
    for index, section in enumerate(exam.sections):
        items = []
        for question in section.questions:
            explanation = question.explanation
            items.append({
                "questionNumber": question.question_number,
                "type": _parser_type(question),
                "text": _plain(question.text) or f"{question.question_number}.",
                "instruction": " ".join(_plain(line) for line in question.instruction),
                "groupContext": " ".join(_plain(line) for line in question.context),
                "options": [_clean_option(option.content) for option in question.options],
                "correctAnswer": _canonical_answer(question),
                "passageIndex": index,
                "explanationWhy": _canonical_explanation_raw(question),
                "evidenceQuote": _plain(explanation.evidence.quote) if explanation else "",
            })
        sections.append({"passage": " ".join(_plain(line) for line in section.passage), "questions": items})
    return {
        "title": _plain(exam.title),
        "type": exam.exam_type,
        "timeLimit": exam.time_limit,
        "audioUrl": exam.audio_url,
        "sections": sections,
    }


def parsed_parser_projection(quiz: Dict[str, Any]) -> Dict[str, Any]:
    if quiz.get("type") == "Writing":
        return {
            "title": _plain(quiz.get("title")),
            "type": "Writing",
            "timeLimit": int(quiz.get("timeLimit") or 0),
            "writingTasks": [
                {
                    "taskNumber": task.get("taskNumber"),
                    "title": _plain(task.get("title")),
                    "instructions": _plain(task.get("instructions")),
                    "minimumWords": task.get("minimumWords"),
                    "recommendedMinutes": task.get("recommendedMinutes"),
                    "mediaUrl": task.get("mediaUrl", ""),
                    "prompt": _plain(task.get("prompt")),
                }
                for task in quiz.get("writingTasks", [])
            ],
        }
    sections = []
    for section in quiz.get("sections", []):
        items = []
        for question in section.get("questions", []):
            manual = question.get("manualExplanation") or {}
            items.append({
                "questionNumber": question.get("questionNumber"),
                "type": question.get("type"),
                "text": _plain(question.get("text")),
                "instruction": _plain(question.get("instruction")),
                "groupContext": _plain(question.get("groupContext")),
                "options": [_clean_option(option) for option in question.get("options", [])],
                "correctAnswer": question.get("correctAnswer"),
                "passageIndex": question.get("passageIndex"),
                "explanationWhy": _plain(manual.get("rawText")),
                "evidenceQuote": _plain((manual.get("parsedQuotes") or [""])[0]),
            })
        sections.append({"passage": _plain(section.get("passage")), "questions": items})
    return {
        "title": _plain(quiz.get("title")),
        "type": quiz.get("type"),
        "timeLimit": int(quiz.get("timeLimit") or 0),
        "audioUrl": quiz.get("audioUrl", ""),
        "sections": sections,
    }


def _differences(expected: Any, actual: Any, path: str = "$") -> List[str]:
    if isinstance(expected, dict) and isinstance(actual, dict):
        output: List[str] = []
        for key in sorted(set(expected) | set(actual)):
            if key not in expected:
                output.append(f"{path}.{key}: unexpected parser value")
            elif key not in actual:
                output.append(f"{path}.{key}: missing parser value")
            else:
                output.extend(_differences(expected[key], actual[key], f"{path}.{key}"))
        return output
    if isinstance(expected, list) and isinstance(actual, list):
        output = []
        if len(expected) != len(actual):
            output.append(f"{path}: expected {len(expected)} items, parser returned {len(actual)}")
        for index, (left, right) in enumerate(zip(expected, actual)):
            output.extend(_differences(left, right, f"{path}[{index}]"))
        return output
    if expected != actual:
        return [f"{path}: expected {expected!r}, parser returned {actual!r}"]
    return []


def _rich_sequence(exam: CanonicalExam) -> List[RichParagraph]:
    """Visible source paragraphs in the exact renderer order (parser tags excluded)."""
    if exam.exam_type == "Writing":
        return [paragraph for task in exam.writing_tasks for paragraph in task.prompt]
    output: List[RichParagraph] = []
    for section in exam.sections:
        output.extend(section.passage)
        seen_groups = set()
        for question in section.questions:
            group_key = (question.block_type, question.group_id or question.question_number)
            first_in_group = group_key not in seen_groups
            seen_groups.add(group_key)
            if first_in_group:
                output.extend(question.instruction)
                if question.block_type == "FLOW_DRAG":
                    output.extend(question.flow_lines or question.context)
                elif question.block_type == "DIAGRAM_LABEL":
                    for box in question.diagram.get("boxes") or []:
                        if isinstance(box, dict):
                            output.extend(RichParagraph.from_value(line) for line in box.get("lines") or [])
                else:
                    output.extend(question.context)
                output.extend(option.content for option in question.options)
            if question.block_type not in {"FLOW_DRAG", "MAP_DRAG", "DIAGRAM_LABEL"}:
                output.append(question.text)
    return [paragraph for paragraph in output if paragraph.text.strip()]


def _native_alignment(paragraph: Any) -> str:
    # python-docx exposes enum values as int-compatible values.
    return {1: "center", 2: "right", 3: "justify"}.get(int(paragraph.alignment or 0), "left")


def verify_native_format_contract(exam: CanonicalExam, path: str | Path) -> None:
    """Verify run-level bold/italic/underline + alignment in the produced DOCX.

    `index.py` verifies semantic parser output below; this native check closes the
    gap before `paragraph_to_html()` flattens Word runs into HTML.
    """
    paragraphs = docx.Document(str(path)).paragraphs
    cursor = 0
    for expected in _rich_sequence(exam):
        expected_text = expected.text
        matched = None
        for index in range(cursor, len(paragraphs)):
            actual_text = paragraphs[index].text
            if actual_text == expected_text or actual_text.endswith(expected_text):
                matched = (index, paragraphs[index])
                break
        if matched is None:
            raise ParserContractError(f"Native DOCX format contract: missing visible paragraph {expected_text!r}.")
        cursor, paragraph = matched
        cursor += 1
        actual_alignment = _native_alignment(paragraph)
        if actual_alignment != expected.alignment:
            raise ParserContractError(
                f"Native DOCX format contract: {expected_text!r} alignment is {actual_alignment}, expected {expected.alignment}."
            )
        runs = paragraph.runs[-len(expected.runs):] if expected.runs else []
        if len(runs) != len(expected.runs):
            raise ParserContractError(f"Native DOCX format contract: run count differs for {expected_text!r}.")
        for actual_run, expected_run in zip(runs, expected.runs):
            if (
                actual_run.text != expected_run.text
                or bool(actual_run.bold) != expected_run.bold
                or bool(actual_run.italic) != expected_run.italic
                or bool(actual_run.underline) != expected_run.underline
            ):
                raise ParserContractError(f"Native DOCX format contract: styled run differs in {expected_text!r}.")
        if expected.list_item and "list" not in (paragraph.style.name or "").lower():
            raise ParserContractError(f"Native DOCX format contract: list style was lost for {expected_text!r}.")


@dataclass
class RoundTripResult:
    quiz: Dict[str, Any]
    expected: Dict[str, Any]
    actual: Dict[str, Any]
    differences: List[str]


def verify_docx_roundtrip(exam: CanonicalExam, path: str | Path) -> RoundTripResult:
    # Direct import is intentional. Do not duplicate parser logic here.
    from api.index import docx_quiz_validation_error, parse_docx_to_quiz

    verify_native_format_contract(exam, path)
    quiz = parse_docx_to_quiz(docx.Document(str(path)))
    parser_error = docx_quiz_validation_error(quiz)
    if parser_error:
        raise ParserContractError(parser_error)
    expected = canonical_parser_projection(exam)
    actual = parsed_parser_projection(quiz)
    differences = _differences(expected, actual)
    if differences:
        preview = "\n".join(differences[:12])
        raise ParserContractError(f"DOCX failed the index.py parser contract:\n{preview}")
    return RoundTripResult(quiz=quiz, expected=expected, actual=actual, differences=[])
