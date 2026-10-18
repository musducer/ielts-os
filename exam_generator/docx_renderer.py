"""Deterministic DOCX renderer for the grammar consumed by ``api/index.py``.

It intentionally emits only the documented control tags.  It never asks an
LLM to format Word content, so a parser-contract failure is reproducible.
"""

from __future__ import annotations

import re
from collections import OrderedDict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Sequence

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Emu, Inches

from .media import ResolvedMediaPaths, encode_media_metadata
from .schema import CanonicalExam, Option, Question, RichParagraph, Section, WritingTask, option_answer_values


class DocxRenderError(ValueError):
    pass


class _MediaRenderContext:
    def __init__(self, exam: CanonicalExam, paths: Dict[str, str]):
        self.exam = exam
        self.paths = paths
        self.resolved = paths if isinstance(paths, ResolvedMediaPaths) else None
        self.rendered: set[str] = set()
        self.rendered_order: List[str] = []
        if self.resolved is not None:
            self.resolved.rendered_occurrence_ids.clear()
            self.occurrences = {item.occurrence_id: item for item in self.resolved.extraction.occurrences}
            self.assets = {item.asset_id: item for item in self.resolved.extraction.assets}
        else:
            self.occurrences = {}
            self.assets = {}

    def payload(self, media: Dict[str, Any]) -> Dict[str, Any]:
        asset_id = str(media.get("asset_id") or "").strip()
        occurrence_id = str(media.get("occurrence_id") or "").strip()
        if not asset_id or not occurrence_id:
            raise DocxRenderError("Rich media needs exact asset_id and occurrence_id before rendering.")
        if self.resolved is None:
            return dict(media)
        occurrence = self.occurrences.get(occurrence_id)
        asset = self.assets.get(asset_id)
        binding = self.resolved.bindings.get(occurrence_id) or {}
        if not occurrence or not asset or occurrence.asset_id != asset_id or not binding:
            raise DocxRenderError(f"Media occurrence {occurrence_id!r} is not bound to the requested asset.")
        result = dict(media)
        result.update({
            "asset_id": asset_id,
            "occurrence_id": occurrence_id,
            "association": str(binding.get("association") or "").strip().lower(),
            "sha256": asset.sha256,
            "width_emu": occurrence.width_emu,
            "height_emu": occurrence.height_emu,
            "alt_text": occurrence.alt_text,
            "title": occurrence.title,
        })
        return result

    def legacy_payload(self, asset_id: str, url: str, association: str) -> Dict[str, Any] | None:
        if not asset_id:
            return None
        if self.resolved is None:
            suffix = re.sub(r"[^A-Za-z0-9_-]", "_", asset_id)
            return {
                "asset_id": asset_id,
                "occurrence_id": f"occ_legacy_{suffix}",
                "association": association,
                "url": url,
            }
        candidates = [
            occurrence_id
            for occurrence_id, binding in self.resolved.bindings.items()
            if str(binding.get("association") or "").strip().lower() == association
            and str(binding.get("asset_id") or "").strip() == asset_id
            and occurrence_id not in self.rendered
        ]
        if not candidates:
            # A rich-media placeholder may already have rendered this dedicated
            # question/task occurrence at its exact source position.
            already_rendered = any(
                str(binding.get("association") or "").strip().lower() == association
                and str(binding.get("asset_id") or "").strip() == asset_id
                and occurrence_id in self.rendered
                for occurrence_id, binding in self.resolved.bindings.items()
            )
            if already_rendered:
                return None
            raise DocxRenderError(f"No bound occurrence is available for {association}.")
        if len(candidates) != 1:
            raise DocxRenderError(
                f"{association} has {len(candidates)} media occurrences but no exact rich-media positions."
            )
        occurrence_id = candidates[0]
        occurrence = self.occurrences[occurrence_id]
        asset = self.assets[asset_id]
        return {
            "asset_id": asset_id,
            "occurrence_id": occurrence_id,
            "association": association,
            "url": url,
            "sha256": asset.sha256,
            "width_emu": occurrence.width_emu,
            "height_emu": occurrence.height_emu,
            "alt_text": occurrence.alt_text,
            "title": occurrence.title,
        }

    def record(self, payload: Dict[str, Any]) -> None:
        occurrence_id = str(payload.get("occurrence_id") or "")
        if occurrence_id in self.rendered:
            raise DocxRenderError(f"Media occurrence {occurrence_id!r} was rendered more than once.")
        self.rendered.add(occurrence_id)
        self.rendered_order.append(occurrence_id)
        if self.resolved is not None:
            self.resolved.rendered_occurrence_ids.append(occurrence_id)

    def finish(self) -> None:
        if self.resolved is None:
            return
        declared = set(self.resolved.bindings)
        if self.rendered != declared:
            missing = sorted(declared - self.rendered)
            extra = sorted(self.rendered - declared)
            raise DocxRenderError(
                f"Media occurrence render contract failed; missing={missing}, unexpected={extra}."
            )


def _plain(value: object) -> str:
    if isinstance(value, RichParagraph):
        value = value.text
    elif isinstance(value, Option):
        value = value.content.text
    elif isinstance(value, dict):
        value = RichParagraph.from_value(value).text
    text = re.sub(r"<[^>]+>", " ", str(value or ""))
    return re.sub(r"\s+", " ", text).strip()


def _safe_url(value: str, field_name: str) -> str:
    value = str(value or "").strip()
    if value and not value.startswith("https://"):
        raise DocxRenderError(f"{field_name} must be a public HTTPS URL.")
    return value


def _picture_size(payload: Dict[str, Any]) -> Dict[str, Emu]:
    width = payload.get("width_emu")
    height = payload.get("height_emu")
    try:
        width = int(width) if width else 0
        height = int(height) if height else 0
    except (TypeError, ValueError):
        width = height = 0
    max_width = int(Inches(6.1))
    if width and height:
        scale = min(1.0, max_width / width)
        return {"width": Emu(round(width * scale)), "height": Emu(round(height * scale))}
    if width:
        return {"width": Emu(min(width, max_width))}
    if height:
        return {"height": Emu(min(height, int(Inches(8.0))))}
    return {"width": Inches(6.1)}


def _embed_media(container: Any, media: Dict[str, Any], context: _MediaRenderContext,
                 alignment: str = "center") -> None:
    payload = context.payload(media)
    asset_id = str(payload.get("asset_id") or "")
    source = context.paths.get(asset_id)
    if not source or not Path(source).is_file():
        raise DocxRenderError(f"Managed media asset {asset_id!r} is unavailable for deterministic DOCX embedding.")
    paragraph = container.add_paragraph()
    paragraph.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }.get(alignment, WD_ALIGN_PARAGRAPH.CENTER)
    try:
        shape = paragraph.add_run().add_picture(source, **_picture_size(payload))
    except Exception as exc:
        raise DocxRenderError(f"Managed media asset {asset_id!r} is not a renderer-safe image.") from exc
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", encode_media_metadata(payload))
    if payload.get("title"):
        doc_pr.set("title", str(payload["title"])[:1_000])
    context.record(payload)


def _add_paragraph(document: Any, value: object, context: _MediaRenderContext, prefix: str = "") -> None:
    paragraph = RichParagraph.from_value(value)
    if paragraph.media:
        if prefix:
            raise DocxRenderError("Media cannot replace a numbered question stem.")
        _embed_media(document, paragraph.media, context, paragraph.alignment)
        return
    if paragraph.table:
        if prefix:
            raise DocxRenderError("A table cannot replace a numbered question stem.")
        rows = paragraph.table["rows"]
        occupied = set()
        placements = []
        columns = 0
        for r, row in enumerate(rows):
            c = 0
            for cell in row["cells"]:
                while (r, c) in occupied:
                    c += 1
                width, height = cell["colspan"], cell["rowspan"]
                if r + height > len(rows):
                    raise DocxRenderError("Table row span extends past the final row.")
                slots = {(rr, cc) for rr in range(r, r + height) for cc in range(c, c + width)}
                if slots & occupied:
                    raise DocxRenderError("Table contains overlapping cell spans.")
                occupied.update(slots)
                placements.append((r, c, cell))
                c += width
                columns = max(columns, c)
        if not columns or columns > 100:
            raise DocxRenderError("Table must have 1-100 columns.")
        target = document.add_table(rows=len(rows), cols=columns)
        target.style = "Table Grid"
        for r, c, cell in placements:
            destination = target.cell(r, c)
            if cell["colspan"] > 1 or cell["rowspan"] > 1:
                destination = destination.merge(target.cell(r + cell["rowspan"] - 1, c + cell["colspan"] - 1))
            destination.text = ""
            for content in cell["paragraphs"]:
                _add_paragraph(destination, content, context)
            # Word cells need a final paragraph; remove only the redundant leading empty one.
            if len(destination.paragraphs) > 1 and not destination.paragraphs[0].text:
                first = destination.paragraphs[0]._p
                first.getparent().remove(first)
        return
    style = "List Bullet" if paragraph.list_item else None
    target = document.add_paragraph(style=style)
    target.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[paragraph.alignment]
    if prefix:
        target.add_run(prefix)
    for source_run in paragraph.runs:
        run = target.add_run(source_run.text)
        run.bold = source_run.bold
        run.italic = source_run.italic
        run.underline = source_run.underline


def _add_lines(document: Any, lines: Iterable[object], context: _MediaRenderContext) -> None:
    for line in lines:
        paragraph = RichParagraph.from_value(line)
        if paragraph.text.strip() or paragraph.media or paragraph.table:
            _add_paragraph(document, paragraph, context)


def _option_label(index: int) -> str:
    if index >= 26:
        raise DocxRenderError("The current parser grammar supports at most 26 labelled options per group.")
    return chr(65 + index)


def _answer_key(question: Question, answer: str) -> str:
    answer = _plain(answer)
    if question.block_type in {"CHOICE", "CHOICE_MULTIPLE"}:
        if re.fullmatch(r"[A-Za-z]", answer):
            return answer.upper()
        for index, option in enumerate(question.options):
            if _plain(option).casefold() == answer.casefold():
                return _option_label(index)
        raise DocxRenderError(
            f"Question {question.question_number}: choice answer {answer!r} is not in its options."
        )
    if question.block_type == "MATCHING":
        if re.fullmatch(r"[A-Za-z]", answer):
            return answer.upper()
        return answer
    return answer


def explanation_text(question: Question) -> str:
    if not question.explanation:
        return ""
    chunks = [_plain(question.explanation.why)]
    evidence = question.explanation.evidence
    if evidence.quote:
        chunks.append(f'"{_plain(evidence.quote)}"')
    if evidence.locator:
        locator = _plain(evidence.locator)
        chunks.append(locator if locator.startswith("[") else f"[{locator}]")
    explanation = question.explanation
    chunks.extend([f"Paraphrase: {_plain(explanation.paraphrase_mapping)}",
                   f"Why correct: {_plain(explanation.correct_reason)}"])
    for option in question.options:
        reason = explanation.distractor_reasons.get(option.id) or explanation.distractor_reasons.get(option.content.text)
        if reason:
            chunks.append(f"{_plain(option)}: {_plain(reason)}")
    chunks.append(f"Takeaway: {_plain(explanation.student_takeaway)}")
    return ' '.join(part for part in chunks if part)


def _write_explanation(document: docx.Document, question: Question) -> None:
    if question.explanation:
        document.add_paragraph(f"[EXPLANATION] {question.question_number}: {explanation_text(question)}")


def _question_groups(section: Section) -> Iterable[List[Question]]:
    """Keep source order while allowing an explicit shared group_id."""
    current_key = None
    current: List[Question] = []
    for question in section.questions:
        key = (question.block_type, question.group_id or f"question-{question.question_number}")
        if current and key != current_key:
            yield current
            current = []
        current_key = key
        current.append(question)
    if current:
        yield current


def _write_option_bank(document: docx.Document, questions: Sequence[Question], context: _MediaRenderContext) -> None:
    leader = questions[0]
    for index, option in enumerate(leader.options):
        _add_paragraph(document, option.content, context, prefix=f"{_option_label(index)}. ")


def _write_generic_group(document: docx.Document, questions: Sequence[Question], context: _MediaRenderContext) -> None:
    leader = questions[0]
    document.add_paragraph(f"[{leader.block_type}]")
    _add_lines(document, leader.instruction, context)
    if leader.left_title:
        document.add_paragraph(f"[LEFT_TITLE] {_plain(leader.left_title)}")
    if leader.right_title:
        document.add_paragraph(f"[RIGHT_TITLE] {_plain(leader.right_title)}")
    if leader.context:
        document.add_paragraph("[CONTEXT]")
        _add_lines(document, leader.context, context)
        document.add_paragraph("[/CONTEXT]")
    if leader.options:
        _write_option_bank(document, questions, context)
    for question in questions:
        _add_paragraph(document, question.text, context, prefix=f"{question.question_number}. ")
        for answer in option_answer_values(question):
            document.add_paragraph(f"*{_answer_key(question, answer)}")
        _write_explanation(document, question)


def _write_flow_group(document: docx.Document, questions: Sequence[Question], context: _MediaRenderContext) -> None:
    leader = questions[0]
    document.add_paragraph("[FLOW_DRAG]")
    _add_lines(document, leader.instruction, context)
    document.add_paragraph("[FLOW]")
    flow_lines = leader.flow_lines or leader.context
    if not flow_lines:
        raise DocxRenderError(f"Question {leader.question_number}: FLOW_DRAG requires flow_lines.")
    _add_lines(document, flow_lines, context)
    document.add_paragraph("[/FLOW]")
    document.add_paragraph("OPTIONS:")
    _write_option_bank(document, questions, context)
    for question in questions:
        document.add_paragraph(f"{question.question_number}.")
        document.add_paragraph(f"*{_answer_key(question, option_answer_values(question)[0])}")
        _write_explanation(document, question)


def _write_map_group(document: docx.Document, questions: Sequence[Question], context: _MediaRenderContext) -> None:
    leader = questions[0]
    media_url = _safe_url(leader.media_url, "MAP_DRAG media_url")
    document.add_paragraph("[MAP_DRAG]")
    _add_lines(document, leader.instruction, context)
    document.add_paragraph(f"IMAGE: {media_url}")
    payload = context.legacy_payload(
        leader.media_asset_id, media_url, f"question:{leader.question_number}"
    )
    if payload:
        _embed_media(document, payload, context)
    document.add_paragraph("[SLOTS]")
    all_slots: Dict[str, Dict[str, float]] = {}
    for question in questions:
        all_slots.update(question.slots)
    for number in sorted(all_slots, key=lambda value: int(value)):
        slot = all_slots[number]
        try:
            x, y = float(slot["x"]), float(slot["y"])
        except (KeyError, TypeError, ValueError) as exc:
            raise DocxRenderError(f"MAP_DRAG slot {number} needs numeric x and y.") from exc
        width, height = float(slot.get("width", 18)), float(slot.get("height", 7))
        document.add_paragraph(f"{number}: x={x}, y={y}, w={width}, h={height}")
    document.add_paragraph("[/SLOTS]")
    document.add_paragraph("OPTIONS:")
    _write_option_bank(document, questions, context)
    for question in questions:
        document.add_paragraph(f"{question.question_number}.")
        document.add_paragraph(f"*{_answer_key(question, option_answer_values(question)[0])}")
        _write_explanation(document, question)


def _write_diagram_group(document: docx.Document, questions: Sequence[Question], context: _MediaRenderContext) -> None:
    leader = questions[0]
    diagram = leader.diagram
    media_url = _safe_url(leader.media_url, "DIAGRAM_LABEL media_url")
    boxes = diagram.get("boxes") or []
    if not isinstance(boxes, list):
        raise DocxRenderError("DIAGRAM_LABEL diagram.boxes must be a list.")
    document.add_paragraph("[DIAGRAM_LABEL]")
    _add_lines(document, leader.instruction, context)
    document.add_paragraph(f"IMAGE: {media_url}")
    payload = context.legacy_payload(
        leader.media_asset_id, media_url, f"question:{leader.question_number}"
    )
    if payload:
        _embed_media(document, payload, context)
    if diagram.get("image_max_width"):
        document.add_paragraph(f"IMAGE_MAX_WIDTH: {int(diagram['image_max_width'])}")
    document.add_paragraph(f"IMAGE_MODE: {_plain(diagram.get('image_mode') or 'TEXT_BOXES')}")
    if diagram.get("image_aspect_ratio"):
        document.add_paragraph(f"IMAGE_ASPECT_RATIO: {_plain(diagram['image_aspect_ratio'])}")
    for box in boxes:
        if not isinstance(box, dict):
            raise DocxRenderError("DIAGRAM_LABEL boxes must be objects.")
        try:
            name = re.sub(r"[^A-Za-z0-9_-]+", "-", str(box["id"]).strip())
            x, y, width, height = (float(box[key]) for key in ("x", "y", "w", "h"))
        except (KeyError, TypeError, ValueError) as exc:
            raise DocxRenderError("Each diagram box needs id, x, y, w, h.") from exc
        document.add_paragraph(f"[BOX {name} x={x} y={y} w={width} h={height}]")
        _add_lines(document, box.get("lines") or [], context)
        document.add_paragraph("[/BOX]")
    document.add_paragraph("[ANSWERS]")
    for question in questions:
        document.add_paragraph(f"{question.question_number}. *{_answer_key(question, option_answer_values(question)[0])}")
        _write_explanation(document, question)
    document.add_paragraph("[/ANSWERS]")
    document.add_paragraph("[/DIAGRAM_LABEL]")


def _render_standard_exam(document: docx.Document, exam: CanonicalExam, context: _MediaRenderContext) -> None:
    document.add_paragraph(f"[TITLE] {_plain(exam.title)}")
    document.add_paragraph(f"[TIME] {exam.time_limit}")
    document.add_paragraph(f"[TYPE] {exam.exam_type}")
    if exam.exam_type == "Listening" and exam.audio_url:
        document.add_paragraph(f"[AUDIO] {_safe_url(exam.audio_url, 'audio_url')}")
    for section in exam.sections:
        document.add_paragraph("[PASSAGE]")
        _add_lines(document, section.passage, context)
        document.add_paragraph("[QUESTIONS]")
        for group in _question_groups(section):
            block_type = group[0].block_type
            if block_type == "FLOW_DRAG":
                _write_flow_group(document, group, context)
            elif block_type == "MAP_DRAG":
                _write_map_group(document, group, context)
            elif block_type == "DIAGRAM_LABEL":
                _write_diagram_group(document, group, context)
            else:
                _write_generic_group(document, group, context)


def _render_writing_exam(document: docx.Document, exam: CanonicalExam, context: _MediaRenderContext) -> None:
    document.add_paragraph(f"[TITLE] {_plain(exam.title)}")
    document.add_paragraph(f"[TIME] {exam.time_limit}")
    document.add_paragraph("[TYPE] Writing")
    for task in sorted(exam.writing_tasks, key=lambda item: item.task_number):
        document.add_paragraph(f"[WRITING_TASK {task.task_number}]")
        if task.title:
            document.add_paragraph(f"[TASK_TITLE] {_plain(task.title)}")
        if task.instructions:
            document.add_paragraph(f"[INSTRUCTIONS] {_plain(task.instructions)}")
        if task.recommended_minutes:
            document.add_paragraph(f"[MINUTES] {task.recommended_minutes}")
        if task.minimum_words:
            document.add_paragraph(f"[MIN_WORDS] {task.minimum_words}")
        document.add_paragraph("[PROMPT]")
        _add_lines(document, task.prompt, context)
        if task.media_url:
            document.add_paragraph(f"[MEDIA] {_safe_url(task.media_url, 'writing task media_url')}")
        payload = context.legacy_payload(
            task.media_asset_id, task.media_url, f"writing_task:{task.task_number}"
        )
        if payload:
            _embed_media(document, payload, context)


def render_docx(exam: CanonicalExam, output_path: str | Path, media_asset_paths: Dict[str, str] | None = None) -> Path:
    """Render an exam with only the current parser's documented grammar."""
    document = docx.Document()
    media_asset_paths = media_asset_paths or {}
    media_context = _MediaRenderContext(exam, media_asset_paths)
    if exam.exam_type == "Writing":
        _render_writing_exam(document, exam, media_context)
    else:
        _render_standard_exam(document, exam, media_context)
    media_context.finish()
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output
